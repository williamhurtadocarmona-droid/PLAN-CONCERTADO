import streamlit as st
import pandas as pd
from docx import Document
import io
import zipfile
import re

st.set_page_config(page_title="Generador de Planes de Trabajo SENA", page_icon="📄", layout="wide")
st.title("📄 Generador de Planes de Trabajo - SENA")

uploaded_excel = st.file_uploader("Cargar Reporte de Juicios Evaluativos (Excel)", type=["xls", "xlsx"])

if uploaded_excel:
    try:
        # 1. Programa y Metadatos
        prog_sel = st.selectbox("Especialidad/Programa:", ["TOTF - Operación en Torno y Fresadora", "TMMI - Mantenimiento Mecánico Industrial"])
        df_meta = pd.read_excel(uploaded_excel, header=None, nrows=12)
        
        num_ficha = next((str(r[2] if pd.notna(r[2]) else r[1]).strip() for _, r in df_meta.iterrows() if str(r[0]).strip() == "Ficha de Caracterización:"), "")
        
        is_totf = "TOTF" in prog_sel
        den_prog = "OPERACION EN TORNO Y FRESADORA" if is_totf else "MECANICA DE MAQUINARIA INDUSTRIAL"
        proj_def = "OPTIMIZACIÓN EN LA FABRICACIÓN DE COMPONENTES MECÁNICOS..." if is_totf else "IMPLEMENTACIÓN DEL PROGRAMA DE MANTENIMIENTO MECÁNICO INDUSTRIAL..."

        st.info(f"📌 **Programa:** {den_prog} | **Ficha:** {num_ficha}")

        # 2. Datos del Proyecto
        c1, c2 = st.columns(2)
        proj_input = c1.text_input("Proyecto Formativo:", value=proj_def)
        fase_input = c2.selectbox("Fase del Proyecto:", ["Análisis", "Planeación", "Ejecución", "Evaluación"], index=2)

        # 3. Selección y Actividades
        st.subheader("🎯 Competencia y Actividades")
        if is_totf:
            st.selectbox("Competencia:", ["290201211 - MECANIZAR PIEZA INDUSTRIAL DE ACUERDO CON TÉCNICAS MANUALES Y SEMIAUTOMÁTICAS"])
            rap_sel = st.selectbox("Resultado de Aprendizaje (RAP):", [
                "694494 - 1.OPERAR TORNO CONVENCIONAL DE ACUERDO CON PROCEDIMIENTOS TÉCNICOS Y NORMATIVA.",
                "694495 - 2.OPERAR FRESADORA CONVENCIONAL DE ACUERDO CON PROCEDIMIENTOS TÉCNICOS Y NORMATIVA"
            ])
            acts_base = [
                "Preparación de herramientas y equipo: Identificar herramientas de corte, afilar y seleccionar elementos de sujeción.",
                "Elaboración de la orden operacional: Analizar plano técnico, secuenciar operaciones y calcular parámetros.",
                "Puesta a punto del torno convencional: Montar pieza, calibrar alturas y pruebas de seguridad.",
                "Ejecución de operaciones de torneado: Cilindrado, refrentado, taladrado y roscado."
            ] if "694494" in rap_sel else [
                "Preparación y puesta a punto de máquina y herramientas de fresado.",
                "Ejecución del fresado según orden operacional y secuencia.",
                "Control dimensional y verificación con instrumentos de metrología.",
                "Identificación de fallas, mejoras y reporte técnico.",
                "Seguridad y protección: Uso de EPP y buenas prácticas."
            ]
        else:
            st.selectbox("Competencia:", ["REPARAR EQUIPOS SEGUN PROCEDIMIENTOS Y MANUALES TECNICOS."])
            rap_sel = st.selectbox("Resultado de Aprendizaje (RAP):", ["OPERAR MÁQUINAS Y HERRAMIENTAS CONVENCIONALES SEGÚN ESPECIFICACIONES TÉCNICAS."])
            acts_base = ["Fabricar elementos mecánicos aplicando procesos de mecanizado con torno convencional."]

        # 4. Estado y Configuración Masiva
        tipo_plan = st.radio("Momento del Plan:", ["Plan Inicial", "Plan Final"], horizontal=True)
        m1, m2 = st.columns(2)
        ent_masiva = m1.selectbox("Forma de Entrega masiva:", ["Personalizar", "Física (Todas)", "Digital (Todas)"])
        est_masiva = m2.selectbox("Estado masivo (Plan Final):", ["Personalizar", "SÍ (Todos)", "NO (Ninguno)"]) if tipo_plan == "Plan Final" else "Personalizar"

        # 5. Detalle de Actividades
        acts_desc, ent_act, est_act = [], [], []
        for i, txt in enumerate(acts_base, 1):
            st.markdown(f"--- **Actividad {i}** ---")
            cols = st.columns(3 if tipo_plan == "Plan Final" else 2)
            acts_desc.append(cols[0].text_area(f"Desc {i}", value=txt, height=75, label_visibility="collapsed"))
            ent_act.append(cols[1].radio(f"Ent {i}", ["Física", "Digital"], index=0 if "Física" in ent_masiva else 1, horizontal=True))
            if tipo_plan == "Plan Final":
                est_act.append(cols[2].radio(f"Est {i}", ["SÍ", "NO"], index=0 if "SÍ" in est_masiva else 1, horizontal=True))

        df_raw = pd.read_excel(uploaded_excel, skiprows=12)
        df_raw.columns = [str(c).strip() for c in df_raw.columns]
        aprendices = df_raw[df_raw['Estado'].str.upper() == 'EN FORMACION'][['Tipo de Documento', 'Número de Documento', 'Nombre', 'Apellidos']].drop_duplicates()

        st.success(f"✅ Aprendices detectados: **{len(aprendices)}**")

        # 6. Generación del ZIP
        if st.button("📦 Generar y Descargar Planes en ZIP (Word Original)", use_container_width=True):
            zip_buffer = io.BytesIO()
            with zipfile.ZipFile(zip_buffer, "a", zipfile.ZIP_DEFLATED, False) as zip_file:
                for _, ap in aprendices.iterrows():
                    doc = Document("Plan de trabajo .docx")
                    
                    reemplazos = {
                        "«Nombre»": str(ap['Nombre']), "«Apellidos»": str(ap['Apellidos']),
                        "«Tipo_de_Doc»": str(ap['Tipo de Documento']), "«N_Documento»": str(ap['Número de Documento']),
                        "TECNICO INSTALACION SISTEMAS ELECTRICOS RESIDENCIALES Y COMERCIALES": den_prog,
                        "«Programa»": den_prog, "«Denominacion»": den_prog,
                        "837101": num_ficha, "«Ficha»": num_ficha, "«Numero_Ficha»": num_ficha,
                        "«Proyecto_Formativo»": proj_input, "«Proyecto»": proj_input,
                        "«Fase_Proyecto»": fase_input, "«Fase»": fase_input
                    }
                    
                    for p in doc.paragraphs:
                        for k, v in reemplazos.items():
                            if k in p.text: p.text = p.text.replace(k, v)
                            
                    if doc.tables:
                        for row in doc.tables[0].rows:
                            for cell in row.cells:
                                for k, v in reemplazos.items():
                                    if k in cell.text: cell.text = cell.text.replace(k, v)
                        if len(doc.tables[0].rows) > 2 and len(doc.tables[0].rows[2].cells) > 5:
                            doc.tables[0].rows[2].cells[6].text = proj_input
                        if len(doc.tables[0].rows) > 3 and len(doc.tables[0].rows[3].cells) > 1:
                            doc.tables[0].rows[3].cells[1].text = fase_input

                    if len(doc.tables) > 1:
                        t_act = doc.tables[1]
                        for idx, desc in enumerate(acts_desc):
                            r_idx = idx + 3
                            if r_idx < len(t_act.rows):
                                cells = t_act.rows[r_idx].cells
                                cells[0].text = rap_sel
                                cells[1].text = str(idx + 1)
                                cells[2].text = desc
                                cells[3].text = "X" if ent_act[idx] == "Física" else ""
                                cells[4].text = "" if ent_act[idx] == "Física" else "X"
                                cells[7].text = "X" if tipo_plan == "Plan Final" and est_act[idx] == "SÍ" else ""
                                cells[8].text = "X" if tipo_plan == "Plan Final" and est_act[idx] == "NO" else ""

                        for r_rm in range(12, len(acts_desc) + 2, -1):
                            if r_rm < len(t_act.rows):
                                tr = t_act.rows[r_rm]._tr
                                tr.getparent().remove(tr)

                    doc_io = io.BytesIO()
                    doc.save(doc_io)
                    sigla = prog_sel.split(' - ')[0]
                    zip_file.writestr(f"Plan_Trabajo_{sigla}_{ap['Número de Documento']}_{ap['Nombre']}_{ap['Apellidos']}.docx", doc_io.getvalue())

            st.download_button("⬇️ Descargar Archivo ZIP", data=zip_buffer.getvalue(), file_name=f"Planes_Trabajo_{prog_sel.split(' - ')[0]}_{tipo_plan.replace(' ', '_')}_Ficha_{num_ficha}.zip", mime="application/zip")

    except Exception as e:
        st.error(f"Error al procesar: {e}")
